import pygame
import os
import openpyxl
from openpyxl.utils import get_column_letter


class Text:
    def __init__(self, screen, text, x, y, fontSize=32, textColor=(0, 0, 0)):
        self.screen = screen
        self.x = x
        self.y = y
        self.fontSize = fontSize
        self.text = text
        self.font = pygame.font.Font(None, fontSize)
        self.textColor = textColor
        self.textSurfaces = self.font.render(self.text, True, self.textColor)

    def draw(self):
        self.screen.blit(self.textSurfaces, (self.x, self.y))

    def setText(self, text):
        self.text = text
        self.textSurfaces = self.font.render(self.text, True, self.textColor)


class Book:
    def __init__(self, id, name, author_name, cover, year):
        if len(id) <= 8:
            if id.find(" "):
                self.id = id
            else:
                print(f"No One space {id}")
                self.id = 0
        else:
            print(f"Too much symbol {id}")
            self.id = 0
        self.name = name
        self.author_name = author_name
        self.cover = cover
        if not any(char.isalpha() for char in year):
            if 100 > int(year) > 26:
                self.year = "19" + year
            elif int(year) <= 26:
                self.year = "20" + year
            elif int(year) > 400:
                self.year = year
        else:
            self.year = 1970
            print(f"Year is not number {id}")

    def print_console(self):
        print(f"Книга номер: {self.id}")
        print(f"  Название: {self.name}")
        print(f"  Автор: {self.author_name}")
        print(f"  Обложка: {self.cover}")
        print(f"  Год: {self.year}")


class ExcelManager:
    def __init__(self, filename="output.xlsx"):
        try:
            self.filename = filename
            if os.path.exists(self.filename):
                self.wb = openpyxl.load_workbook(self.filename)
                self.ws = self.wb.active
            else:
                self.wb = openpyxl.Workbook()
                self.ws = self.wb.active
        except Exception as e:
            print("Error!")

    def _resize_column(self, column_idx):
        try:
            col_letter = get_column_letter(column_idx)
            max_len = 0
            for row in range(1, self.ws.max_row + 1):
                cell_value = self.ws.cell(row=row, column=column_idx).value
                if cell_value is not None:
                    max_len = max(max_len, len(str(cell_value)))
            self.ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        except Exception as e:
            print("Error!")

    def write_cell(self, cell_coord, value):
        self.ws[cell_coord] = value
        column_idx = self.ws[cell_coord].column
        self._resize_column(column_idx)

    def append_row(self, row_values):
        self.ws.append(row_values)
        for col_idx in range(1, len(row_values) + 1):
            self._resize_column(col_idx)

    def add_book_with_safeguards(self, book, row_num):
        self.write_cell(f"A{row_num}", row_num-4)
        try:
            book_id_parts = book.id.split()
            formatted_id = f"{book_id_parts[0]}-({book_id_parts[1]})"
            self.write_cell(f"B{row_num}", formatted_id)
        except Exception:
            self.write_cell(f"B{row_num}", "11111-(1)")

        self.write_cell(f"C{row_num}", book.author_name)
        self.write_cell(f"D{row_num}", book.name)

        try:
            self.write_cell(f"E{row_num}", book.year)
        except Exception:
            self.write_cell(f"E{row_num}", 1970)

        self.write_cell(f"F{row_num}", 1)
        self.write_cell(f"G{row_num}", book.cover)
        self.write_cell(f"H{row_num}", "потрепаная, выпадают страницы")

    def save(self):
        self.wb.save(self.filename)
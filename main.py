import pygame_gui
from openpyxl import Workbook
from openpyxl import load_workbook

from classes import *

pygame.init()
size = (800, 600)
screen = pygame.display.set_mode(size)
clock = pygame.time.Clock()
books = []
input_path_manager = pygame_gui.UIManager(size)
input_data = pygame_gui.UIManager(size)

stages = ["input_path_file", 'add_new_book']
stage_number = 0
stage = stages[stage_number]

input_path_text = pygame_gui.elements.UITextEntryLine(
    relative_rect=pygame.Rect((330, 6), (350, 32)),
    manager=input_path_manager
)
input_path_button = pygame_gui.elements.UIButton(
    relative_rect=pygame.Rect((690, 5), (48, 32)),
    text='Далее',
    manager=input_path_manager
)

input_inventory_num_text = pygame_gui.elements.UITextEntryLine(
    relative_rect=pygame.Rect((250, 6), (240, 32)),
    manager=input_data
)
input_name_text = pygame_gui.elements.UITextEntryLine(
    relative_rect=pygame.Rect((250, 46), (240, 32)),
    manager=input_data
)
input_author_name_text = pygame_gui.elements.UITextEntryLine(
    relative_rect=pygame.Rect((250, 86), (240, 32)),
    manager=input_data
)
input_year_text = pygame_gui.elements.UITextEntryLine(
    relative_rect=pygame.Rect((250, 126), (240, 32)),
    manager=input_data
)

input_cover_button = pygame_gui.elements.UIDropDownMenu(
    options_list=["Тверд", "Мягкий"],
    starting_option="Тверд",
    relative_rect=pygame.Rect((620, 6), (100, 32)),
    manager=input_data
)

input_data_button = pygame_gui.elements.UIButton(
    relative_rect=pygame.Rect((653, 554), (128, 32)),
    text='Добавить книгу',
    manager=input_data
)

ws = None
wb = None
last_string_num = 0
excel_manager = None

running = True
while running:
    time_delta = clock.tick(60) / 1000.0
    for event in pygame.event.get():
        if event.type == pygame.QUIT:
            running = False
        elif event.type == pygame_gui.UI_BUTTON_PRESSED:
            if stage == "input_path_file":
                if event.ui_element == input_path_button:
                    stage_number += 1
                    stage = stages[stage_number]
                    wb = load_workbook(input_path_text.get_text())
                    ws = wb.active
                    tmp = 1
                    while ws[f"B{tmp}"].value is not None:
                        tmp += 1
                    last_string_num = tmp
                    excel_manager = ExcelManager(input_path_text.get_text())
            elif stage == "add_new_book":
                if event.ui_element == input_data_button:
                    # print(input_cover_button.selected_option[0])
                    books.append(Book(input_inventory_num_text.get_text(), input_name_text.get_text(),
                                      input_author_name_text.get_text(), input_cover_button.selected_option[0],
                                      input_year_text.get_text()))
                    books[-1].print_console()
                    input_inventory_num_text.set_text("")
                    input_name_text.set_text("")
                    input_author_name_text.set_text("")
                    input_year_text.set_text("")

                    excel_manager.add_book_with_safeguards(books[-1], last_string_num)
                    excel_manager.save()

                    last_string_num += 1
        if stage == "input_path_file":
            input_path_manager.process_events(event)
        elif stage == "add_new_book":
            input_data.process_events(event)
    screen.fill((255, 255, 255))
    if stage == "input_path_file":
        Text(screen, "Введите путь к файлу excel:", 10, 10, fontSize=32).draw()
        input_path_manager.update(time_delta)
        input_path_manager.draw_ui(screen)
    elif stage == "add_new_book":
        Text(screen, "Инвентарный номер:", 10, 10).draw()
        Text(screen, "Обложка:", 500, 10).draw()
        Text(screen, "Название книги:", 10, 50).draw()
        Text(screen, "Автор книги:", 10, 90).draw()
        Text(screen, "Год издания:", 10, 130).draw()
        input_data.update(time_delta)
        input_data.draw_ui(screen)
    pygame.display.flip()
    clock.tick(60)
pygame.quit()
